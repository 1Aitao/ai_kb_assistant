"""文档内容提取器：按扩展名分发到对应提取函数，新增格式只需在 EXTRACTORS 加一行。"""
import logging
from functools import lru_cache

import cv2
import pymupdf as fitz
import numpy as np
from pypdf import PdfReader

from config import settings

logger = logging.getLogger(__name__)


@lru_cache
def _get_ocr():
    """懒加载全局唯一的 RapidOCR 实例：模型只加载一次，重复上传不再重复初始化。

    加载后立即用空白图做一次预热推理：ONNX Runtime 首次推理有较大
    惰性初始化开销（内存池/线程池），预热把它挪到服务启动时，
    避免服务重启后的第一次上传额外慢约 50%。
    """
    from rapidocr_onnxruntime import RapidOCR
    ocr = RapidOCR()
    ocr(np.zeros((64, 64, 3), dtype=np.uint8), use_cls=False)
    return ocr


def warm_up_ocr() -> None:
    """预热 OCR 引擎（加载模型 + 空白图推理），在服务启动时调用可消除首次上传的冷启动延迟。"""
    _get_ocr()


def extract_txt_md(file_path: str) -> str:
    """提取 txt/md 文本文件内容。"""
    with open(file_path, "r", encoding="utf-8") as f:
        return f.read()


def _ocr_page(ocr, page) -> str:
    """OCR 识别单个 PDF 页：渲染成图像后在内存中直接识别，不写临时文件。"""
    # 渲染倍率可配置（默认 1.2）：实测比 2.0 快约 28% 且识别质量相当
    scale = settings.OCR_RENDER_SCALE
    pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale))
    # 直接在内存中传递图像字节，避免 Windows 上临时文件被占用导致删除失败
    img = cv2.imdecode(np.frombuffer(pix.tobytes("png"), np.uint8), cv2.IMREAD_COLOR)
    # use_cls=False 跳过文本方向分类：扫描件几乎都是正向的，跳过可省推理时间
    result, _ = ocr(img, use_cls=False)
    if not result:
        return ""
    return "\n".join(line[1] for line in result)


def extract_pdf(file_path: str) -> str:
    """提取 PDF 文字：逐页优先读文字层，只有无文字层的页（扫描页）才降级 OCR。

    逐页判断的好处：混合型 PDF（部分文字页 + 部分扫描页）只 OCR 扫描页；
    纯文字 PDF 完全不加载 OCR 模型。
    """
    reader = PdfReader(file_path)
    page_texts = [(p.extract_text() or "").strip() for p in reader.pages]

    # 全部页都有文字层，无需 OCR
    if all(page_texts):
        return "\n".join(page_texts).strip()

    ocr = _get_ocr()
    doc = fitz.open(file_path)
    parts = []
    total = len(page_texts)
    for idx, text in enumerate(page_texts):
        if text:
            parts.append(text)
            continue
        logger.info("OCR 识别第 %d/%d 页 ...", idx + 1, total)
        parts.append(_ocr_page(ocr, doc[idx]))
    return "\n".join(parts).strip()


def extract_image(file_path: str) -> str:
    """提取图片（png/jpg）中的文字。"""
    result, _ = _get_ocr()(file_path, use_cls=False)
    if not result:
        return ""
    return "\n".join(line[1] for line in result).strip()


def extract_docx(file_path: str) -> str:
    """提取 Word 文档段落文字（类名重命名避免与业务模型冲突）。"""
    from docx import Document as DocxDocument
    d = DocxDocument(file_path)
    return "\n".join(p.text for p in d.paragraphs if p.text.strip())


def extract_xlsx(file_path: str) -> str:
    """提取 Excel：read_only+data_only 取单元格值，每行用 | 拼接。"""
    from openpyxl import load_workbook
    wb = load_workbook(file_path, read_only=True, data_only=True)
    lines = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


# 扩展名 -> 提取函数 的路由表
EXTRACTORS = {
    ".txt": extract_txt_md, ".md": extract_txt_md,
    ".pdf": extract_pdf,
    ".png": extract_image, ".jpg": extract_image, ".jpeg": extract_image,
    ".docx": extract_docx,
    ".xlsx": extract_xlsx,
}
